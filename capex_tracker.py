"""
AI Compute Capex Tracker
========================
Pulls quarterly capex, revenue, operating cash flow and depreciation for any
US-listed company from the SEC XBRL API, normalizes everything to calendar
quarters, and computes the metrics used to judge AI infrastructure spending:
growth, capex intensity, free cash flow, and capex versus depreciation.

Pipeline:  fetch  ->  extract  ->  clean / derive quarters  ->  metrics  ->  export

Usage:
    python capex_tracker.py                    # default: the five hyperscalers
    python capex_tracker.py NVDA AVGO          # any tickers you like
    python capex_tracker.py MSFT --refresh     # re-download latest data from SEC
"""

import argparse
import json
import time
from pathlib import Path

import pandas as pd
import requests

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------

# SEC requires a User-Agent with your name and email, or requests get blocked.
USER_AGENT = "Your Name your.email@example.com"  # change to a placeholder before uploading to GitHub

# Tracked when you run the script without naming any tickers.
DEFAULT_TICKERS = ["MSFT", "GOOGL", "AMZN", "META", "ORCL"]

# Each metric: (label, XBRL tags in priority order). Companies label the same
# line item differently; if two tags cover the same quarter, the earlier wins.
# If a company shows "missing" for a metric, find the tag it uses in its 10-K
# and add it to the list.
METRICS = {
    "capex": ("Capex", [
        "PaymentsToAcquirePropertyPlantAndEquipment",
        "PaymentsToAcquireProductiveAssets",
        "PaymentsForCapitalImprovements",
    ]),
    "revenue": ("Revenue", [
        "RevenueFromContractWithCustomerExcludingAssessedTax",
        "Revenues",
        "SalesRevenueNet",
        "RevenueFromContractWithCustomerIncludingAssessedTax",
    ]),
    "ocf": ("Operating cash flow", [
        "NetCashProvidedByUsedInOperatingActivities",
        "NetCashProvidedByUsedInOperatingActivitiesContinuingOperations",
    ]),
    "da": ("Depreciation & amortization", [
        "DepreciationDepletionAndAmortization",
        "DepreciationAmortizationAndAccretionNet",
        "DepreciationAndAmortization",
        "DepreciationAmortizationAndOther",
        "Depreciation",
    ]),
}
CAPEX_TAGS = METRICS["capex"][1]  # kept for backwards compatibility

START_YEAR = 2020  # first calendar year to keep in the output

BASE_DIR = Path(__file__).parent
RAW_DIR = BASE_DIR / "data" / "raw"
OUT_DIR = BASE_DIR / "output"
OVERRIDES_FILE = BASE_DIR / "manual_overrides.csv"
API_URL = "https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json"
TICKER_URL = "https://www.sec.gov/files/company_tickers.json"

# Chart colors: a fixed categorical order, validated for colorblind separation.
SERIES_COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4",
                 "#008300", "#4a3aa7", "#e34948"]


# ---------------------------------------------------------------------------
# 1. FETCH
# ---------------------------------------------------------------------------

def lookup_companies(tickers, refresh=False):
    """Turn tickers into (CIK, company name) using the SEC's official ticker list."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    cache = RAW_DIR / "company_tickers.json"
    if cache.exists() and not refresh:
        data = json.loads(cache.read_text())
    else:
        resp = requests.get(TICKER_URL, headers={"User-Agent": USER_AGENT}, timeout=30)
        resp.raise_for_status()
        cache.write_text(resp.text)
        data = resp.json()

    by_ticker = {v["ticker"].upper(): (int(v["cik_str"]), v["title"]) for v in data.values()}
    found = {}
    for t in tickers:
        t = t.upper()
        if t in by_ticker:
            found[t] = by_ticker[t]
        else:
            print(f"  {t}: ticker not found in SEC list, skipped")
    return found


def fetch_company_facts(ticker, cik, refresh=False):
    """Download a company's full XBRL fact set, caching it locally as JSON."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    cache = RAW_DIR / f"{ticker}.json"
    if cache.exists() and not refresh:
        return json.loads(cache.read_text())

    resp = requests.get(API_URL.format(cik=cik),
                        headers={"User-Agent": USER_AGENT}, timeout=30)
    resp.raise_for_status()
    cache.write_text(resp.text)
    time.sleep(0.2)  # stay well under SEC's 10 requests/second limit
    return resp.json()


# ---------------------------------------------------------------------------
# 2. EXTRACT
# ---------------------------------------------------------------------------

def extract_facts(ticker, facts, tags):
    """Pull every fact for the given tags from 10-Q / 10-K filings."""
    us_gaap = facts.get("facts", {}).get("us-gaap", {})
    rows = []
    for priority, tag in enumerate(tags):
        for f in us_gaap.get(tag, {}).get("units", {}).get("USD", []):
            if "start" not in f or not f.get("form", "").startswith(("10-Q", "10-K")):
                continue
            rows.append({
                "ticker": ticker, "tag": tag, "tag_priority": priority,
                "start": f["start"], "end": f["end"], "value": f["val"],
                "form": f["form"], "filed": f["filed"], "accn": f["accn"],
            })
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    for col in ("start", "end", "filed"):
        df[col] = pd.to_datetime(df[col])
    return df


def extract_capex_facts(ticker, facts):
    return extract_facts(ticker, facts, CAPEX_TAGS)


# ---------------------------------------------------------------------------
# 3. CLEAN: dedupe, then turn year-to-date figures into single quarters
# ---------------------------------------------------------------------------

def dedupe(df):
    """
    The same period appears in many filings (e.g. last year's figure is repeated
    as a comparative). Keep one row per period: preferred tag first, then the
    most recently filed value, which picks up any restatement.
    """
    df = df.sort_values(["tag_priority", "filed"], ascending=[True, False])
    return df.drop_duplicates(subset=["start", "end"], keep="first")


def calendar_quarter(start, end):
    """Label a period by the calendar quarter containing its midpoint.
    Handles odd fiscal years: Oracle's Jun-Aug quarter becomes Q3."""
    mid = start + (end - start) / 2
    return f"{mid.year}Q{(mid.month - 1) // 3 + 1}"


def derive_quarters(df):
    """
    Cash flow statements in 10-Qs are year-to-date: Q2 shows 6 months, Q3 shows
    9 months, the 10-K shows 12 months. Within each fiscal year (same start
    date), subtract the previous YTD figure to get the single quarter.
    Income statement items also report 3-month figures directly; those are
    picked up as "reported" quarters.
    """
    out = []
    for (ticker, start), grp in df.groupby(["ticker", "start"]):
        prev_val, q_start = 0, start
        for _, row in grp.sort_values("end").iterrows():
            q_val = row["value"] - prev_val
            days = (row["end"] - q_start).days + 1
            if 80 <= days <= 100:  # a clean single quarter
                out.append({
                    "ticker": ticker,
                    "cal_quarter": calendar_quarter(q_start, row["end"]),
                    "period_start": q_start.date(),
                    "period_end": row["end"].date(),
                    "capex_usd": q_val,
                    "method": "reported" if prev_val == 0 else "derived (YTD diff)",
                    "tag": row["tag"], "tag_priority": row.get("tag_priority", 0),
                    "form": row["form"], "accn": row["accn"],
                })
            # if a quarter is missing (gap > 100 days), skip it but keep going
            prev_val, q_start = row["value"], row["end"] + pd.Timedelta(days=1)

    q = pd.DataFrame(out)
    if q.empty:
        return q
    # a quarter can be reported directly AND derived; keep one
    q = q.sort_values("method", ascending=False)  # "reported" before "derived"
    q = q.drop_duplicates(subset=["ticker", "period_end"], keep="first")
    q = q[q["cal_quarter"].str[:4].astype(int) >= START_YEAR]
    return q.sort_values(["ticker", "period_end"]).reset_index(drop=True)


def build_metric(ticker, facts, metric):
    """
    Derive single quarters for one metric. Each tag is processed on its own, so
    a year-to-date subtraction never mixes two different line items; then, for
    each quarter, the highest-priority tag wins.
    """
    raw = extract_facts(ticker, facts, METRICS[metric][1])
    if raw.empty:
        return pd.DataFrame()
    pieces = [derive_quarters(dedupe(grp)) for _, grp in raw.groupby("tag")]
    pieces = [p for p in pieces if not p.empty]
    if not pieces:
        return pd.DataFrame()
    q = pd.concat(pieces, ignore_index=True).sort_values("tag_priority")
    q = q.drop_duplicates(subset=["ticker", "cal_quarter"], keep="first")
    q = q.rename(columns={"capex_usd": "value_usd"}).drop(columns="tag_priority")
    q.insert(1, "metric", metric)
    return q.sort_values("cal_quarter").reset_index(drop=True)


def apply_overrides(q, ticker, metric="capex"):
    """
    Some quarters are missing from the SEC API (e.g. the company used its own
    custom XBRL tag that period). Fill them by hand in manual_overrides.csv,
    copied from the filing, with the source noted. A manual value replaces any
    derived value for the same quarter.

    Accepts either column layout:
      ticker,cal_quarter,capex_usd_millions,source            (capex only)
      ticker,cal_quarter,metric,value_usd_millions,source     (any metric)
    """
    if not OVERRIDES_FILE.exists():
        return q
    ov = pd.read_csv(OVERRIDES_FILE)
    if ov.empty:
        return q
    if "metric" not in ov.columns:
        ov["metric"] = "capex"
    value_col = "value_usd_millions" if "value_usd_millions" in ov.columns else "capex_usd_millions"
    ov = ov[(ov["ticker"].str.upper() == ticker) & (ov["metric"].str.lower() == metric)]
    if ov.empty:
        return q
    value_name = "value_usd" if "value_usd" in q.columns else "capex_usd"
    rows = pd.DataFrame({
        "ticker": ticker,
        "cal_quarter": ov["cal_quarter"].values,
        "period_start": None, "period_end": None,
        value_name: ov[value_col].astype(float).values * 1e6,
        "method": "manual (from filing)",
        "tag": "manual", "form": ov["source"].values, "accn": "",
    })
    if "metric" in q.columns:
        rows.insert(1, "metric", metric)
    if not q.empty:
        q = q[~q["cal_quarter"].isin(rows["cal_quarter"])]
    return pd.concat([q, rows], ignore_index=True).sort_values("cal_quarter").reset_index(drop=True)


# ---------------------------------------------------------------------------
# 4. METRICS
# ---------------------------------------------------------------------------

def quarter_range(first, last):
    """All calendar quarters between two labels, so gaps become blank rows."""
    p = pd.period_range(pd.Period(first, freq="Q"), pd.Period(last, freq="Q"), freq="Q")
    return [str(x) for x in p]


def company_panel(long_df, ticker):
    """Rows = every calendar quarter (gaps kept as blanks), columns = metrics in $bn."""
    d = long_df[long_df["ticker"] == ticker]
    panel = d.pivot_table(index="cal_quarter", columns="metric",
                          values="value_usd", aggfunc="sum") / 1e9
    panel = panel.reindex(columns=list(METRICS))
    capex_q = d.loc[d["metric"] == "capex", "cal_quarter"]
    first, last = capex_q.min(), capex_q.max()
    return panel.reindex(quarter_range(first, last))


def compute_metrics(panel):
    """
    The same calculations the Excel sheets do with formulas, in pandas, for the
    charts and the written summary. Gaps stay blank rather than being guessed.
    """
    p = panel.copy()
    c, r, o, da = p["capex"], p["revenue"], p["ocf"], p["da"]
    ttm = lambda s: s.rolling(4, min_periods=4).sum()
    p["capex_qoq"] = c / c.shift(1) - 1
    p["capex_yoy"] = c / c.shift(4) - 1
    p["capex_ttm"] = ttm(c)
    p["capex_intensity"] = c / r
    p["fcf"] = o - c
    p["fcf_ttm"] = ttm(p["fcf"])
    p["fcf_margin_ttm"] = p["fcf_ttm"] / ttm(r)
    p["capex_to_da"] = c / da
    yoy = p["capex_yoy"]
    p["trend"] = [
        "" if pd.isna(a) or pd.isna(b) else ("Accelerating" if a > b else "Decelerating")
        for a, b in zip(yoy, yoy.shift(1))
    ]
    return p


def to_wide(long_df, tickers, metric="capex"):
    """Rows = calendar quarter, columns = company, values in $bn."""
    d = long_df[long_df["metric"] == metric] if "metric" in long_df.columns else long_df
    value_col = "value_usd" if "value_usd" in d.columns else "capex_usd"
    wide = d.pivot_table(index="cal_quarter", columns="ticker",
                         values=value_col, aggfunc="sum") / 1e9
    return wide[[t for t in tickers if t in wide.columns]].sort_index()


def trim_incomplete_tail(wide):
    """
    Companies report on different schedules (Oracle's quarter lands a month
    before the others). Drop the latest quarters until every company has
    reported, so totals are never understated. Earlier gaps, e.g. a company
    that IPO'd recently, stay as blanks instead of cutting the whole table.
    """
    complete = wide.notna().all(axis=1)
    if not complete.any():
        return wide
    return wide.loc[:complete[complete].index[-1]]


# ---------------------------------------------------------------------------
# 5. EXPORT
# ---------------------------------------------------------------------------

# Columns on each company sheet. Inputs are blue (data from filings), the rest
# are Excel formulas so the workbook recalculates if an input is corrected.
COMPANY_COLS = [
    ("Calendar Quarter", None),
    ("Capex ($bn)", "input"), ("Revenue ($bn)", "input"),
    ("Operating CF ($bn)", "input"), ("D&A ($bn)", "input"),
    ("Capex QoQ", "0.0%"), ("Capex YoY", "0.0%"), ("Capex TTM ($bn)", "#,##0.0"),
    ("Capex / Revenue", "0.0%"), ("FCF ($bn)", "#,##0.0;(#,##0.0)"),
    ("FCF TTM ($bn)", "#,##0.0;(#,##0.0)"), ("FCF Margin TTM", "0.0%;(0.0%)"),
    ("Capex / D&A", "0.00x"), ("Capex Trend", None),
]


def company_row_formulas(i):
    """Formulas for row i of a company sheet (row 2 is the first quarter)."""
    f = {}
    f["F"] = f'=IF(OR(B{i}="",B{i-1}=""),"",B{i}/B{i-1}-1)' if i >= 3 else None
    f["G"] = f'=IF(OR(B{i}="",B{i-4}=""),"",B{i}/B{i-4}-1)' if i >= 6 else None
    f["H"] = f'=IF(COUNT(B{i-3}:B{i})<4,"",SUM(B{i-3}:B{i}))' if i >= 5 else None
    f["I"] = f'=IF(OR(B{i}="",C{i}="",C{i}=0),"",B{i}/C{i})'
    f["J"] = f'=IF(OR(B{i}="",D{i}=""),"",D{i}-B{i})'
    f["K"] = f'=IF(COUNT(J{i-3}:J{i})<4,"",SUM(J{i-3}:J{i}))' if i >= 5 else None
    f["L"] = (f'=IF(OR(K{i}="",COUNT(C{i-3}:C{i})<4),"",K{i}/SUM(C{i-3}:C{i}))'
              if i >= 5 else None)
    f["M"] = f'=IF(OR(B{i}="",E{i}="",E{i}=0),"",B{i}/E{i})'
    f["N"] = (f'=IF(OR(G{i}="",G{i-1}=""),"",IF(G{i}>G{i-1},"Accelerating","Decelerating"))'
              if i >= 7 else None)
    return f


def export_excel(long_df, panels, wide, companies, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    font, bold = Font(name="Arial"), Font(name="Arial", bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    blue = Font(name="Arial", color="0000FF")    # hardcoded data
    green = Font(name="Arial", color="008000")   # links to another sheet
    tickers = list(wide.columns)
    latest = wide.index[-1]
    k = len(tickers)

    def header(sheet, widths=16):
        for c in sheet[1]:
            c.font, c.fill = bold, header_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "B2"
        for col_idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = widths

    # --- Company sheets (built first so other sheets can link to them) -----
    row_of = {}  # (ticker, quarter) -> row number on that company's sheet
    company_sheets = []
    for t in tickers:
        ws = wb.create_sheet(t)
        company_sheets.append(ws)
        ws.append([name for name, _ in COMPANY_COLS])
        panel = panels[t]
        for i, (quarter, vals) in enumerate(panel.iterrows(), start=2):
            row_of[(t, quarter)] = i
            ws.cell(i, 1, quarter).font = font
            for j, m in enumerate(METRICS, start=2):
                v = vals[m]
                c = ws.cell(i, j, None if pd.isna(v) else float(v))
                c.font, c.number_format = blue, "#,##0.0"
            for col, formula in company_row_formulas(i).items():
                j = ord(col) - 64
                c = ws.cell(i, j, formula)
                c.font = font
                fmt = COMPANY_COLS[j - 1][1]
                if fmt and fmt != "input":
                    c.number_format = fmt
        header(ws, 14)

    # --- Summary: snapshot of the latest quarter every company has reported -
    summ = wb.create_sheet("Summary", 0)
    summ.append([f"Snapshot: {latest}", "Capex ($bn)", "Capex YoY", "Capex TTM ($bn)",
                 "Share of Group Capex", "Capex / Revenue", "FCF TTM ($bn)",
                 "FCF Margin TTM", "Capex / D&A", "Capex Trend"])
    src_cols = ["B", "G", "H", None, "I", "K", "L", "M", "N"]
    fmts = ["#,##0.0", "0.0%", "#,##0.0", "0.0%", "0.0%", "#,##0.0;(#,##0.0)",
            "0.0%;(0.0%)", "0.00x", None]
    first, last = 2, k + 1
    for i, t in enumerate(tickers, start=2):
        summ.cell(i, 1, f"{t} ({companies[t][1]})").font = font
        r = row_of[(t, latest)]
        for j, (src, fmt) in enumerate(zip(src_cols, fmts), start=2):
            if src is None:  # share of group capex
                formula = f"=IFERROR(B{i}/SUM($B${first}:$B${last}),\"\")"
                c = summ.cell(i, j, formula)
                c.font = font
            else:
                c = summ.cell(i, j, f"='{t}'!{src}{r}")
                c.font = green
            if fmt:
                c.number_format = fmt
    tot = k + 2
    summ.cell(tot, 1, "Group total").font = bold
    c = summ.cell(tot, 2, f"=SUM(B{first}:B{last})"); c.font, c.number_format = bold, "#,##0.0"
    c = summ.cell(tot, 4, f"=SUM(D{first}:D{last})"); c.font, c.number_format = bold, "#,##0.0"
    c = summ.cell(tot, 7, f"=SUM(G{first}:G{last})"); c.font, c.number_format = bold, "#,##0.0;(#,##0.0)"
    summ.cell(tot + 2, 1, "Green cells link to each company's sheet; black cells are formulas. "
                          "See Notes for definitions.").font = font
    header(summ, 16)
    summ.column_dimensions["A"].width = 34

    # --- Quarterly Capex across companies, with total and share ------------
    qc = wb.create_sheet("Quarterly Capex", 1)
    qc.append(["Calendar Quarter"] + [f"{t} ($bn)" for t in tickers] + ["Total ($bn)"]
              + [f"{t} Share" for t in tickers])
    total_col = get_column_letter(k + 2)
    for i, quarter in enumerate(wide.index, start=2):
        qc.cell(i, 1, quarter).font = font
        for j, t in enumerate(tickers, start=2):
            r = row_of.get((t, quarter))
            c = qc.cell(i, j, f"=IF('{t}'!B{r}=\"\",\"\",'{t}'!B{r})" if r else None)
            c.font, c.number_format = green, "#,##0.0"
        c = qc.cell(i, k + 2, f"=SUM(B{i}:{get_column_letter(k + 1)}{i})")
        c.font, c.number_format = bold, "#,##0.0"
        for j in range(k):
            col = get_column_letter(j + 2)
            c = qc.cell(i, k + 3 + j, f'=IF(OR({col}{i}="",{total_col}{i}=0),"",{col}{i}/{total_col}{i})')
            c.font, c.number_format = font, "0.0%"
    header(qc, 14)

    # --- YoY growth across companies, all formulas --------------------------
    yoy = wb.create_sheet("YoY Growth", 2)
    yoy.append(["Calendar Quarter"] + [f"{t} YoY" for t in tickers] + ["Total YoY"])
    for i in range(2, len(wide) + 2):
        yoy.cell(i, 1, f"='Quarterly Capex'!A{i}").font = font
        for j in range(2, k + 3):
            col = get_column_letter(j)
            f = (f"=IFERROR('Quarterly Capex'!{col}{i}/'Quarterly Capex'!{col}{i - 4}-1,\"\")"
                 if i >= 6 else None)
            c = yoy.cell(i, j, f)
            c.font, c.number_format = font, "0.0%"
    header(yoy, 14)

    # --- Audit trail: every input traced back to its filing ----------------
    audit = wb.create_sheet("Audit Trail")
    cols = ["ticker", "metric", "cal_quarter", "period_start", "period_end", "value_usd",
            "method", "tag", "form", "accn"]
    audit.append(cols)
    for _, r in long_df[cols].iterrows():
        audit.append([str(v) if k_ in ("period_start", "period_end") and v is not None else v
                      for k_, v in r.items()])
    for row in audit.iter_rows(min_row=2):
        for c in row:
            c.font = font
        row[5].number_format = "#,##0"
    header(audit, 16)

    # --- Notes ----------------------------------------------------------------
    notes = wb.create_sheet("Notes")
    lines = [
        "Methodology",
        "Source: SEC XBRL companyfacts API (data.sec.gov), 10-Q and 10-K filings.",
        "10-Q cash flows are year-to-date; single quarters are derived by subtracting the prior YTD figure.",
        "Fiscal quarters are mapped to the calendar quarter containing the period midpoint "
        "(MSFT FY ends June, ORCL FY ends May).",
        "Summary and cross-company sheets use the latest quarter every company has reported; "
        "company sheets show all quarters available.",
        "",
        "Definitions",
        "Capex: cash paid for property and equipment (cash flow statement).",
        "Capex / Revenue (capex intensity): share of each revenue dollar reinvested in infrastructure.",
        "FCF: operating cash flow minus capex. FCF TTM: sum of the last four quarters.",
        "FCF Margin TTM: FCF TTM divided by revenue TTM.",
        "Capex / D&A: above 1x means the asset base is growing, so depreciation will keep "
        "rising and weigh on future operating margins.",
        "Capex Trend: Accelerating if capex YoY growth is higher than the previous quarter's.",
        "",
        "Tags used (priority order)",
    ] + [f"{label}: " + ", ".join(tags) for label, tags in METRICS.values()] + [
        "",
        "Caveats",
        "Cash capex excludes assets acquired under finance leases, which are material for some "
        "companies (e.g. Microsoft, Amazon); company-reported total capex can be higher.",
        "D&A definitions differ by company (some include amortization of intangibles, content, "
        "or lease assets), so compare Capex / D&A trends within a company more than across companies.",
        "Revenue is total company revenue; cloud segment revenue is not in this API.",
        "Blank cells mean the figure was not available in the API; fill them in manual_overrides.csv.",
        "Blue = data from filings; green = link to another sheet; black = formula.",
    ]
    for i, line in enumerate(lines, start=1):
        is_heading = line in ("Methodology", "Definitions", "Tags used (priority order)", "Caveats")
        notes.cell(i, 1, line).font = bold if is_heading else font
    notes.column_dimensions["A"].width = 130

    del wb["Sheet"]
    wb.save(path)


def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.size": 10, "axes.spines.top": False, "axes.spines.right": False,
        "axes.edgecolor": "#8a8984", "axes.labelcolor": "#52514e",
        "xtick.color": "#52514e", "ytick.color": "#52514e",
        "axes.titlecolor": "#0b0b0b", "axes.titleweight": "bold",
    })
    return plt


def export_chart(wide, names, path):
    plt = _plt()
    fig, ax = plt.subplots(figsize=(12, 6))
    colors = SERIES_COLORS[: len(wide.columns)]
    wide.fillna(0).plot(kind="bar", stacked=True, ax=ax, width=0.8, color=colors,
                        edgecolor="white", linewidth=1)
    ax.set_title("Quarterly Capex: " + ", ".join(names), loc="left")
    ax.set_ylabel("Capex ($bn, cash basis)")
    ax.set_xlabel("")
    if len(wide.columns) > 1:
        ax.legend(title="", frameon=False)
    else:
        ax.get_legend().remove()
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    plt.xticks(rotation=45, ha="right")
    plt.tight_layout()
    fig.savefig(path, dpi=150, facecolor="white")
    plt.close(fig)


def export_metrics_chart(metrics, tickers, path, since=None):
    """Small multiples: one panel per metric, so each keeps its own scale."""
    plt = _plt()
    import matplotlib.ticker as mtick
    panels = [
        ("capex_yoy", "Capex YoY growth", mtick.PercentFormatter(1.0, decimals=0), 0),
        ("capex_intensity", "Capex / Revenue", mtick.PercentFormatter(1.0, decimals=0), None),
        ("fcf_margin_ttm", "FCF margin (TTM)", mtick.PercentFormatter(1.0, decimals=0), 0),
        ("capex_to_da", "Capex / D&A", mtick.FormatStrFormatter("%.1fx"), 1),
    ]
    fig, axes = plt.subplots(2, 2, figsize=(13, 8))
    for ax, (col, title, fmt, ref) in zip(axes.flat, panels):
        for color, t in zip(SERIES_COLORS, tickers):
            s = metrics[t][col]
            if since:
                s = s[s.index >= since]
            ax.plot(s.index, s.values, color=color, linewidth=2, label=t, marker="o", markersize=3)
        if ref is not None:
            ax.axhline(ref, color="#8a8984", linewidth=1, linestyle="--")
        ax.set_title(title, loc="left")
        ax.yaxis.set_major_formatter(fmt)
        ax.grid(axis="y", alpha=0.25)
        labels = ax.get_xticks()
        ax.set_xticks(labels[::2])
        ax.tick_params(axis="x", rotation=45)
    handles, labels = axes.flat[0].get_legend_handles_labels()
    if len(tickers) > 1:
        fig.legend(handles, labels, loc="upper center", ncol=len(tickers), frameon=False,
                   bbox_to_anchor=(0.5, 1.0))
    fig.suptitle("")
    plt.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, dpi=150, facecolor="white")
    plt.close(fig)


def fmt_pct(x):
    return "n/a" if pd.isna(x) else f"{x:+.0%}" if abs(x) < 10 else f"{x:+.1f}x"


def build_findings(metrics, wide, companies):
    """Plain-language observations, computed from the numbers (no judgement calls)."""
    tickers = list(wide.columns)
    latest = wide.index[-1]
    snap = pd.DataFrame({t: metrics[t].loc[latest] for t in tickers}).T
    total = wide.sum(axis=1)
    out = []
    if len(total) >= 5 and total.iloc[-5] > 0:
        out.append(f"Group capex in {latest} was ${total.iloc[-1]:.1f}bn, "
                   f"{fmt_pct(total.iloc[-1] / total.iloc[-5] - 1)} year over year.")
    if len(tickers) > 1:
        yoy = snap["capex_yoy"].dropna().astype(float)
        if len(yoy) > 1:
            out.append(f"Fastest capex growth: {yoy.idxmax()} ({fmt_pct(yoy.max())} YoY); "
                       f"slowest: {yoy.idxmin()} ({fmt_pct(yoy.min())}).")
        share = (snap["capex"] / snap["capex"].sum()).astype(float)
        out.append(f"Largest spender: {share.idxmax()}, {share.max():.0%} of group capex.")
    inten = snap["capex_intensity"].dropna().astype(float)
    if len(inten):
        out.append(f"Highest capex intensity: {inten.idxmax()}, spending "
                   f"{inten.max():.0%} of revenue on capex this quarter.")
    fcf = snap["fcf_ttm"].dropna().astype(float)
    neg = list(fcf[fcf < 0].index)
    if neg:
        out.append("Negative free cash flow over the last four quarters: " + ", ".join(neg) + ".")
    fm = snap["fcf_margin_ttm"].dropna().astype(float)
    if len(fm):
        out.append(f"Lowest TTM FCF margin: {fm.idxmin()} ({fm.min():.0%}).")
    cda = snap["capex_to_da"].dropna().astype(float)
    if len(cda):
        out.append(f"Highest capex / D&A: {cda.idxmax()} ({cda.max():.1f}x), "
                   "so its depreciation charge is set to keep rising.")
    acc = [t for t in tickers if snap.loc[t, "trend"] == "Accelerating"]
    dec = [t for t in tickers if snap.loc[t, "trend"] == "Decelerating"]
    if acc or dec:
        out.append("Capex growth accelerating: " + (", ".join(acc) or "none")
                   + "; decelerating: " + (", ".join(dec) or "none") + ".")
    for t in tickers:
        newer = metrics[t].index[metrics[t].index > latest]
        newer = [q for q in newer if not pd.isna(metrics[t].loc[q, "capex"])]
        if newer:
            q = newer[-1]
            out.append(f"{t} has already reported {q}: capex ${metrics[t].loc[q, 'capex']:.1f}bn "
                       f"({fmt_pct(metrics[t].loc[q, 'capex_yoy'])} YoY), "
                       "not yet included in group totals.")
    missing = [f"{t} {METRICS[m][0]}" for t in tickers for m in ("revenue", "ocf", "da")
               if pd.isna(snap.loc[t, m])]
    if missing:
        out.append("Missing in the latest quarter (check tags or add overrides): "
                   + ", ".join(missing) + ".")
    return out, snap


def export_markdown(metrics, wide, companies, path):
    """A summary with computed findings, ready to drop into the LLM Wiki."""
    findings, snap = build_findings(metrics, wide, companies)
    latest = wide.index[-1]
    total = wide.sum(axis=1)
    lines = [f"# AI Capex Tracker: data through {latest}", "",
             "Source: SEC XBRL API (10-Q/10-K). Generated by capex_tracker.py. "
             "Figures in $bn; cash capex excludes finance leases.", "",
             "## Key findings", ""]
    lines += [f"- {f}" for f in findings]
    lines += ["", f"## Snapshot: {latest}", "",
              "| Company | Capex | YoY | Capex/Revenue | FCF TTM | FCF margin TTM | Capex/D&A | Trend |",
              "|---|---|---|---|---|---|---|---|"]
    num = lambda v, f: "" if pd.isna(v) else f.format(v)
    for t, r in snap.iterrows():
        lines.append(f"| {t} | {num(r['capex'], '{:.1f}')} | {num(r['capex_yoy'], '{:+.0%}')} | "
                     f"{num(r['capex_intensity'], '{:.0%}')} | {num(r['fcf_ttm'], '{:.1f}')} | "
                     f"{num(r['fcf_margin_ttm'], '{:.0%}')} | {num(r['capex_to_da'], '{:.1f}x')} | "
                     f"{r['trend']} |")
    lines += ["", "## Quarterly capex, last 8 quarters", "",
              "| Quarter | " + " | ".join(wide.columns) + " | Total |",
              "|---" * (len(wide.columns) + 2) + "|"]
    for quarter in wide.index[-8:]:
        vals = ["" if pd.isna(v) else f"{v:.1f}" for v in wide.loc[quarter]]
        lines.append(f"| {quarter} | " + " | ".join(vals) + f" | {total[quarter]:.1f} |")
    lines += ["", "Numbers have not been checked against the filings unless noted. "
                  "Verify before citing."]
    path.write_text("\n".join(lines) + "\n")
    return findings


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def run(tickers=None, refresh=False):
    companies = lookup_companies(tickers or DEFAULT_TICKERS, refresh=refresh)
    # each set of tickers gets its own folder, so runs don't overwrite each other
    out_dir = OUT_DIR / ("hyperscalers" if not tickers else "_".join(companies))
    out_dir.mkdir(parents=True, exist_ok=True)

    frames = []
    for ticker, (cik, name) in list(companies.items()):
        facts = fetch_company_facts(ticker, cik, refresh=refresh)
        status = []
        for metric in METRICS:
            q = apply_overrides(build_metric(ticker, facts, metric), ticker, metric)
            if metric == "capex" and q.empty:
                break
            status.append(f"{metric} {len(q)}q" if not q.empty else f"{metric} MISSING")
            if not q.empty:
                frames.append(q)
        if not status:
            print(f"  {ticker}: no capex tag found, check METRICS['capex']")
            companies.pop(ticker)
            continue
        print(f"  {ticker} ({name}): " + ", ".join(status))

    if not frames:
        print("No data found for any ticker.")
        return None, None, None
    long_df = pd.concat(frames, ignore_index=True)
    tickers_found = [t for t in companies if t in set(long_df["ticker"])]

    panels = {t: company_panel(long_df, t) for t in tickers_found}
    metrics = {t: compute_metrics(p) for t, p in panels.items()}
    wide = trim_incomplete_tail(to_wide(long_df, tickers_found))

    long_df.to_csv(out_dir / "capex_quarterly.csv", index=False)
    pd.concat({t: m for t, m in metrics.items()}, names=["ticker", "cal_quarter"]).to_csv(
        out_dir / "metrics.csv")
    export_excel(long_df, panels, wide, companies, out_dir / "capex_tracker.xlsx")
    export_chart(wide, [companies[t][1] for t in wide.columns], out_dir / "capex_trend.png")
    since = str(pd.Period(wide.index[-1], freq="Q") - 15)  # last four years
    export_metrics_chart(metrics, list(wide.columns), out_dir / "metrics_dashboard.png", since)
    findings = export_markdown(metrics, wide, companies, out_dir / "capex_summary.md")
    print("\nKey findings:")
    for f in findings:
        print(f"  - {f}")
    print(f"\nDone. Outputs in {out_dir}/")
    return long_df, wide, metrics


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AI compute capex tracker")
    parser.add_argument("tickers", nargs="*", help="e.g. NVDA AVGO (default: five hyperscalers)")
    parser.add_argument("--refresh", action="store_true", help="re-download latest data from SEC")
    args = parser.parse_args()
    run(tickers=args.tickers, refresh=args.refresh)
